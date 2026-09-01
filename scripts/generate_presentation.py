#!/usr/bin/env python3
"""Generate professional AI Travel Inspirator hackathon presentation."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
OUTPUT = ROOT / "docs" / "AI_Travel_Inspirator.pptx"
ARCH_DIR = ROOT / "docs" / "architecture"
AGENT_PNG = ARCH_DIR / "agent_architecture.png"
DEMO_PNG = ARCH_DIR / "demo_vs_future.png"

# ── Brand palette ─────────────────────────────────────────────────────────────
NAVY = RGBColor(0x0A, 0x0E, 0x1A)
NAVY_MID = RGBColor(0x12, 0x18, 0x28)
CARD = RGBColor(0x1C, 0x24, 0x38)
CARD_LIGHT = RGBColor(0x24, 0x2E, 0x45)
GOLD = RGBColor(0xD4, 0xA5, 0x6A)
GOLD_LIGHT = RGBColor(0xE8, 0xC9, 0x9A)
TEAL = RGBColor(0x4E, 0xC4, 0xB0)
CORAL = RGBColor(0xE8, 0x7A, 0x6A)
SKY = RGBColor(0x7B, 0xB8, 0xE8)
WHITE = RGBColor(0xFA, 0xFA, 0xF8)
SILVER = RGBColor(0xB0, 0xB8, 0xC8)
MUTED = RGBColor(0x7A, 0x82, 0x96)
GREEN = RGBColor(0x6B, 0xC4, 0x9A)

FONT = "Calibri"
FONT_LIGHT = "Calibri Light"
FONT_MONO = "Consolas"
TEAM = "dot_thinker"
ORG = "Coforge"
SLIDE_NUM = [0]


def _n() -> int:
    SLIDE_NUM[0] += 1
    return SLIDE_NUM[0]


def _font(p, text, size=14, bold=False, color=WHITE, name=FONT, align=None):
    p.text = text
    p.font.name = name
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    if align:
        p.alignment = align


def new_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = NAVY
    return slide


def add_bg_decor(slide, accent=GOLD):
    """Subtle geometric accents — no circles or rings."""
    wash = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(6.5), Inches(0), Inches(3.5), Inches(2.2))
    wash.fill.solid()
    wash.fill.fore_color.rgb = accent
    wash.fill.transparency = 0.92
    wash.line.fill.background()

    bar = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(0.06), Inches(7.05))
    bar.fill.solid()
    bar.fill.fore_color.rgb = accent
    bar.line.fill.background()

    line = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, Inches(7.35), Inches(10), Inches(0.04))
    line.fill.solid()
    line.fill.fore_color.rgb = accent
    line.line.fill.background()


def add_footer(slide, num: int, section: str = ""):
    bar = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, Inches(7.05), Inches(10), Inches(0.32))
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY_MID
    bar.line.fill.background()

    left = slide.shapes.add_textbox(Inches(0.4), Inches(7.08), Inches(4), Inches(0.25))
    _font(left.text_frame.paragraphs[0], f"{TEAM}  ·  {ORG}", 9, color=MUTED)

    if section:
        mid = slide.shapes.add_textbox(Inches(3.5), Inches(7.08), Inches(3), Inches(0.25))
        _font(mid.text_frame.paragraphs[0], section.upper(), 9, color=SILVER, align=PP_ALIGN.CENTER)

    right = slide.shapes.add_textbox(Inches(9.0), Inches(7.08), Inches(0.6), Inches(0.25))
    _font(right.text_frame.paragraphs[0], str(num), 9, bold=True, color=GOLD, align=PP_ALIGN.RIGHT)


def add_header(slide, title: str, subtitle: str = "", accent=GOLD):
    strip = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, Inches(10), Inches(1.15))
    strip.fill.solid()
    strip.fill.fore_color.rgb = NAVY_MID
    strip.line.fill.background()

    bar = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, Inches(0.1), Inches(1.15))
    bar.fill.solid()
    bar.fill.fore_color.rgb = accent
    bar.line.fill.background()

    tb = slide.shapes.add_textbox(Inches(0.45), Inches(0.2), Inches(9), Inches(0.55))
    _font(tb.text_frame.paragraphs[0], title, 26, bold=True, color=WHITE, name=FONT)

    if subtitle:
        sb = slide.shapes.add_textbox(Inches(0.45), Inches(0.72), Inches(9), Inches(0.35))
        _font(sb.text_frame.paragraphs[0], subtitle, 13, color=SILVER, name=FONT_LIGHT)


def add_card(slide, x, y, w, h, fill=CARD, border=GOLD, radius=True):
    shape_type = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    card = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    card.fill.solid()
    card.fill.fore_color.rgb = fill
    card.line.color.rgb = border
    card.line.width = Pt(0.75)
    return card


def add_bullets_pro(slide, items, left=0.55, top=1.35, width=8.9, size=14, accent=GOLD):
    y = top
    for item in items:
        dot = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(left), Inches(y + 0.07), Inches(0.08), Inches(0.08))
        dot.fill.solid()
        dot.fill.fore_color.rgb = accent
        dot.line.fill.background()

        tb = slide.shapes.add_textbox(Inches(left + 0.25), Inches(y), Inches(width - 0.25), Inches(0.55))
        tf = tb.text_frame
        tf.word_wrap = True
        _font(tf.paragraphs[0], item, size, color=WHITE)
        y += 0.62 if len(item) > 80 else 0.52


def add_callout_pro(slide, text, y=6.0, accent=GOLD):
    add_card(slide, 0.5, y, 9.0, 0.75, CARD_LIGHT, accent)
    tb = slide.shapes.add_textbox(Inches(0.7), Inches(y + 0.15), Inches(8.6), Inches(0.45))
    _font(tb.text_frame.paragraphs[0], text, 13, color=accent, align=PP_ALIGN.CENTER)


def add_box_text(slide, x, y, w, h, text, fill=CARD, text_color=WHITE, border=TEAL, size=11, bold=False):
    card = add_card(slide, x, y, w, h, fill, border)
    tf = card.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    _font(tf.paragraphs[0], text, size, bold, text_color, align=PP_ALIGN.CENTER)
    return card


def add_arrow(slide, x1, y1, x2, y2, color=TEAL):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    conn.line.color.rgb = color
    conn.line.width = Pt(1.75)


def try_add_picture(slide, path: Path, left, top, width, height):
    if path.exists():
        slide.shapes.add_picture(str(path), Inches(left), Inches(top), Inches(width), Inches(height))
        return True
    return False


def slide_section(prs, number: str, title: str, subtitle: str, accent=GOLD):
    slide = new_slide(prs)
    add_bg_decor(slide, accent)
    num = _n()

    big = slide.shapes.add_textbox(Inches(0.8), Inches(2.4), Inches(2), Inches(1))
    _font(big.text_frame.paragraphs[0], number, 52, bold=True, color=accent, name=FONT_LIGHT)

    line = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.8), Inches(3.5), Inches(1.8), Inches(0.05))
    line.fill.solid()
    line.fill.fore_color.rgb = accent
    line.line.fill.background()

    tb = slide.shapes.add_textbox(Inches(0.8), Inches(3.75), Inches(8.5), Inches(0.9))
    _font(tb.text_frame.paragraphs[0], title, 36, bold=True, color=WHITE)

    sb = slide.shapes.add_textbox(Inches(0.8), Inches(4.65), Inches(8), Inches(0.6))
    _font(sb.text_frame.paragraphs[0], subtitle, 16, color=SILVER, name=FONT_LIGHT)

    add_footer(slide, num)


# ── Slides ────────────────────────────────────────────────────────────────────

def slide_title(prs):
    slide = new_slide(prs)
    add_bg_decor(slide, GOLD)
    num = _n()

    badge = add_card(slide, 0.75, 0.55, 2.6, 0.42, NAVY_MID, TEAL)
    _font(badge.text_frame.paragraphs[0], "HACKATHON 2026", 9, bold=True, color=TEAL, align=PP_ALIGN.CENTER)

    coforge = slide.shapes.add_textbox(Inches(7.5), Inches(0.62), Inches(2), Inches(0.3))
    _font(coforge.text_frame.paragraphs[0], ORG, 10, color=MUTED, align=PP_ALIGN.RIGHT)

    tb = slide.shapes.add_textbox(Inches(0.75), Inches(1.85), Inches(8.5), Inches(1.1))
    _font(tb.text_frame.paragraphs[0], "AI Travel Inspirator", 48, bold=True, color=GOLD)

    sub = slide.shapes.add_textbox(Inches(0.75), Inches(3.05), Inches(8.5), Inches(0.55))
    _font(sub.text_frame.paragraphs[0], "Emotion-First Global Travel Discovery Agent", 22, color=WHITE, name=FONT_LIGHT)

    gold_line = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.75), Inches(3.75), Inches(2.5), Inches(0.04))
    gold_line.fill.solid()
    gold_line.fill.fore_color.rgb = GOLD
    gold_line.line.fill.background()

    tag = slide.shapes.add_textbox(Inches(0.75), Inches(4.0), Inches(8.2), Inches(0.8))
    _font(
        tag.text_frame.paragraphs[0],
        "Helping undecided travelers discover destinations that match how they feel —\n"
        "bridging social inspiration and OTA booking with emotion, context, and explainable boards.",
        14,
        color=SILVER,
        name=FONT_LIGHT,
    )

    pills = ["Multi-Modal Emotion", "3 Ps Context", "Multi-Currency", "Live Demo"]
    px = 0.75
    for pill in pills:
        p = add_card(slide, px, 5.15, 2.05, 0.38, CARD, TEAL)
        _font(p.text_frame.paragraphs[0], pill, 9, bold=True, color=TEAL, align=PP_ALIGN.CENTER)
        px += 2.2

    team_line = slide.shapes.add_textbox(Inches(0.75), Inches(5.85), Inches(8), Inches(0.35))
    _font(team_line.text_frame.paragraphs[0], f"Presented by Team {TEAM}", 12, color=MUTED)

    add_footer(slide, num)


def slide_team(prs):
    slide = new_slide(prs)
    add_bg_decor(slide, TEAL)
    add_header(slide, "Meet the Team", f"Team {TEAM}  ·  {ORG}", TEAL)
    num = _n()

    banner = add_card(slide, 3.0, 1.35, 4.0, 0.9, CARD_LIGHT, GOLD)
    tf = banner.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    _font(tf.paragraphs[0], TEAM, 30, bold=True, color=GOLD, align=PP_ALIGN.CENTER)
    p2 = tf.add_paragraph()
    _font(p2, ORG, 11, color=SILVER, align=PP_ALIGN.CENTER)

    members = [
        ("Priyanka Singh", "94872", "Priyanka.4.S@coforge.com", "PS", GOLD, 0.65),
        ("Avinash Tripathi", "76765", "avinash.tripathi@coforge.com", "AT", TEAL, 5.35),
    ]

    for name, emp, email, initials, color, x in members:
        add_card(slide, x, 2.55, 4.0, 3.35, CARD, color)

        strip = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x), Inches(2.55), Inches(4.0), Inches(0.12))
        strip.fill.solid()
        strip.fill.fore_color.rgb = color
        strip.line.fill.background()

        av = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
            Inches(x + 1.45), Inches(2.95), Inches(1.1), Inches(1.1),
        )
        av.fill.solid()
        av.fill.fore_color.rgb = color
        av.line.fill.background()
        av_tb = slide.shapes.add_textbox(Inches(x + 1.45), Inches(3.15), Inches(1.1), Inches(0.65))
        _font(av_tb.text_frame.paragraphs[0], initials, 26, bold=True, color=NAVY, align=PP_ALIGN.CENTER)

        n_box = slide.shapes.add_textbox(Inches(x + 0.2), Inches(4.2), Inches(3.6), Inches(0.4))
        _font(n_box.text_frame.paragraphs[0], name, 18, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

        role = slide.shapes.add_textbox(Inches(x + 0.2), Inches(4.6), Inches(3.6), Inches(0.3))
        _font(role.text_frame.paragraphs[0], "Team Member", 10, color=MUTED, align=PP_ALIGN.CENTER)

        e_box = slide.shapes.add_textbox(Inches(x + 0.2), Inches(4.95), Inches(3.6), Inches(0.3))
        _font(e_box.text_frame.paragraphs[0], f"Emp ID: {emp}", 11, color=SILVER, align=PP_ALIGN.CENTER)

        m_box = slide.shapes.add_textbox(Inches(x + 0.2), Inches(5.3), Inches(3.6), Inches(0.35))
        _font(m_box.text_frame.paragraphs[0], email, 10, color=color, align=PP_ALIGN.CENTER)

    add_callout_pro(slide, "Project: AI Travel Inspirator — Emotion-first global travel discovery powered by AI", 6.05, TEAL)
    add_footer(slide, num, "Team")


def slide_executive_summary(prs):
    slide = new_slide(prs)
    add_bg_decor(slide)
    add_header(slide, "Executive Summary", "Quantitative opportunity · Gap · Goal")
    num = _n()

    stats = [
        ("$622.6B→$1.44T", "Online travel 2025→2034\n~9.75% CAGR (IMARC)", GOLD),
        ("70%+ / ~63%", "Bookings online /\nmobile bookings", TEAL),
        ("80%+", "Travelers research\nvia social media (2025)", CORAL),
        ("APAC 31.8%", "Share of online travel\n(IMARC)", SKY),
    ]
    sx = 0.45
    for label, val, col in stats:
        add_card(slide, sx, 1.3, 2.25, 1.25, CARD, col)
        tb = slide.shapes.add_textbox(Inches(sx + 0.1), Inches(1.4), Inches(2.05), Inches(1.05))
        _font(tb.text_frame.paragraphs[0], label, 12, bold=True, color=col, align=PP_ALIGN.CENTER)
        p2 = tb.text_frame.add_paragraph()
        _font(p2, val, 9, color=WHITE, align=PP_ALIGN.CENTER)
        sx += 2.35

    add_bullets_pro(
        slide,
        [
            "OTA market ~USD 943B (2025, TBRC) / ~USD 561B OTA segment 2026 (Mordor — alternate estimate).",
            "Huge booking market exists — but discovery/emotion layer between social inspiration and OTA checkout is underserved.",
            "Top research sources: search 46%, reviews 36%, friends/family 35%, OTAs 28% — emotion rarely enters the loop.",
            "Our agent: multi-modal emotion + 3 Ps context + explainable inspiration boards — global, multi-currency.",
        ],
        top=2.75,
        size=12,
    )
    add_callout_pro(
        slide,
        "Frame: Capture intent emotionally → enrich with Profile/Preferences/Policy → inspire before travelers hit price filters.",
        y=5.95,
    )
    add_footer(slide, num, "Overview")


def slide_market_opportunity(prs):
    slide = new_slide(prs)
    add_bg_decor(slide, GOLD)
    add_header(slide, "Global Market Opportunity", "Industry estimates — online travel & OTA scale", GOLD)
    num = _n()

    metrics = [
        ("Online Travel", "USD 622.6B (2025)", "→ USD 1.44T by 2034", "~9.75% CAGR · IMARC", GOLD),
        ("OTA Market", "~USD 943B (2025)", "TBRC estimate", "Alt: ~USD 561B OTA '26 (Mordor)", TEAL),
        ("Channel Mix", "70%+ online bookings", "~63% mobile bookings", "Discovery still fragmented", CORAL),
        ("Social Research", "80%+ use social", "to research destinations", "Inspiration ≠ personalization", SKY),
    ]
    mx = 0.45
    for title, a, b, c, col in metrics:
        add_card(slide, mx, 1.3, 2.25, 2.55, CARD, col)
        strip = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(mx), Inches(1.3), Inches(2.25), Inches(0.1))
        strip.fill.solid()
        strip.fill.fore_color.rgb = col
        strip.line.fill.background()
        tb = slide.shapes.add_textbox(Inches(mx + 0.1), Inches(1.5), Inches(2.05), Inches(2.2))
        _font(tb.text_frame.paragraphs[0], title, 11, bold=True, color=col, align=PP_ALIGN.CENTER)
        for line in (a, b, c):
            p = tb.text_frame.add_paragraph()
            _font(p, line, 10, color=WHITE if line == a else SILVER, align=PP_ALIGN.CENTER)
        mx += 2.35

    add_box_text(
        slide, 0.45, 4.05, 9.1, 0.55,
        "Asia-Pacific ~31.8% of online travel (IMARC) — emotions and aspirations cross every market border",
        CARD_LIGHT, GOLD, GOLD, 11, True,
    )

    sources = [
        ("Search engines", "46%"),
        ("Review sites", "36%"),
        ("Friends / family", "35%"),
        ("OTAs", "28%"),
    ]
    sx = 0.45
    for label, pct in sources:
        add_card(slide, sx, 4.8, 2.25, 0.85, CARD, MUTED)
        tb = slide.shapes.add_textbox(Inches(sx + 0.1), Inches(4.9), Inches(2.05), Inches(0.65))
        _font(tb.text_frame.paragraphs[0], pct, 16, bold=True, color=TEAL, align=PP_ALIGN.CENTER)
        p = tb.text_frame.add_paragraph()
        _font(p, label, 10, color=SILVER, align=PP_ALIGN.CENTER)
        sx += 2.35

    add_callout_pro(
        slide,
        "Opportunity wedge: social inspires · OTAs transact · nobody owns emotion-aware, policy-aware discovery.",
        y=5.85,
        accent=TEAL,
    )
    add_footer(slide, num, "Market")


def slide_problem(prs):
    slide = new_slide(prs)
    add_bg_decor(slide, CORAL)
    add_header(slide, "Problem Statement", "Why travelers struggle — globally, not just one market", CORAL)
    num = _n()

    problems = [
        ("Decision Paralysis", "Aspirational travelers drown in options with no personal, emotional guide.", CORAL),
        ("Popularity Bias", "Trending destinations ≠ individual mood, culture, or travel personality.", GOLD),
        ("Currency & Cost Ambiguity", "Budgets span USD / EUR / GBP / INR / AED / JPY / AUD — tools rarely unify guidance.", TEAL),
        ("Emotion Ignored", "OTAs optimise price & availability; social inspires without cost/policy context.", SKY),
        ("Context Gap", "No Profile · Preferences · Policy (3 Ps) layer between inspiration and booking.", CORAL),
    ]
    y = 1.35
    for title, desc, col in problems:
        add_card(slide, 0.55, y, 0.08, 0.5, col, col, radius=False)
        tb = slide.shapes.add_textbox(Inches(0.75), Inches(y), Inches(8.7), Inches(0.5))
        _font(tb.text_frame.paragraphs[0], title, 13, bold=True, color=col)
        p2 = tb.text_frame.add_paragraph()
        _font(p2, desc, 11, color=SILVER)
        y += 0.58

    add_callout_pro(
        slide,
        "Insight: Travelers need a recommendation partner between Instagram and Booking.com — not another top-10 list.",
        accent=CORAL,
    )
    add_footer(slide, num, "Problem")


def slide_solution_proposal(prs):
    slide = new_slide(prs)
    add_bg_decor(slide, TEAL)
    add_header(slide, "Solution Proposal", "Emotion + 3 Ps context + explainable inspiration", TEAL)
    num = _n()

    add_box_text(
        slide, 0.55, 1.3, 8.9, 0.5,
        "AI Travel Inspirator — Multi-modal, context-aware travel discovery agent (global)",
        CARD_LIGHT, GOLD, GOLD, 12, True,
    )

    phases = [
        ("01", "Sense", "Mood chips · selfie\nvoice · text sentiment", GOLD),
        ("02", "Context", "Profile · Preferences\nPolicy (3 Ps)", TEAL),
        ("03", "Reason", "TravelAgent engines\n+ LLM (Gemini)", CORAL),
        ("04", "Inspire", "Boards · scores\nmulti-currency costs", SKY),
    ]
    px = 0.55
    for num_str, title, desc, col in phases:
        add_card(slide, px, 2.0, 2.05, 1.7, CARD, col)
        nb = slide.shapes.add_textbox(Inches(px + 0.15), Inches(2.1), Inches(0.5), Inches(0.35))
        _font(nb.text_frame.paragraphs[0], num_str, 18, bold=True, color=col)
        tb = slide.shapes.add_textbox(Inches(px + 0.15), Inches(2.5), Inches(1.75), Inches(1.1))
        _font(tb.text_frame.paragraphs[0], title, 13, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        p2 = tb.text_frame.add_paragraph()
        _font(p2, desc, 10, color=SILVER, align=PP_ALIGN.CENTER)
        if px < 6.5:
            add_arrow(slide, px + 2.1, 2.85, px + 2.45, 2.85, col)
        px += 2.25

    add_bullets_pro(
        slide,
        [
            "Design principle: Emotions first, logistics second — every recommendation explains emotional fit.",
            "Context beyond single-shot LLM: 3 Ps enrich prompts; RAG + social graph on the roadmap.",
            "Global by default: multi-currency cost guidance (USD · EUR · GBP · INR · AED · JPY · AUD).",
        ],
        top=4.0,
        size=12,
    )
    add_footer(slide, num, "Solution")


def slide_agentic_architecture_overview(prs):
    slide = new_slide(prs)
    add_bg_decor(slide)
    add_header(slide, "Agentic Architecture", "Emotion Layer · Context (3 Ps) · RAG · TravelAgent · LLM")
    num = _n()

    layers = [
        ("Emotion", "Manual mood (DEMO) · selfie CV · voice emotion · text sentiment", GOLD, 1.3),
        ("Context — 3 Ps", "Profile · Preferences · Policy  (+ social/friend graph — FUTURE)", TEAL, 2.15),
        ("RAG Knowledge", "Destination KB · reviews · culture/seasonality — FUTURE retrieve", CORAL, 3.0),
        ("TravelAgent", "Emotional Engine · Destination Matcher · Inspiration Cluster", SKY, 3.85),
        ("LLM + Output", "Gemini 2.5 (+ mock) → concepts, boards, multi-currency costs", GREEN, 4.7),
    ]
    for label, desc, col, y in layers:
        add_card(slide, 0.5, y, 1.85, 0.7, CARD_LIGHT, col)
        lb = slide.shapes.add_textbox(Inches(0.6), Inches(y + 0.15), Inches(1.65), Inches(0.4))
        _font(lb.text_frame.paragraphs[0], label, 10, bold=True, color=col, align=PP_ALIGN.CENTER)
        db = slide.shapes.add_textbox(Inches(2.5), Inches(y + 0.15), Inches(6.95), Inches(0.45))
        _font(db.text_frame.paragraphs[0], desc, 12, color=WHITE)

    add_callout_pro(
        slide,
        "3 Ps enrich context beyond single-shot prompting — personalisation that respects budget, style, and constraints.",
        y=5.7,
        accent=TEAL,
    )
    add_footer(slide, num, "Architecture")


def slide_three_ps_and_emotion(prs):
    slide = new_slide(prs)
    add_bg_decor(slide, TEAL)
    add_header(slide, "3 Ps Context + Multi-Modal Emotion", "What enriches the agent beyond a bare LLM call", TEAL)
    num = _n()

    for title, items, col, x in [
        ("Profile", ["Identity & travel history", "Personality archetype", "Party composition", "Accessibility needs"], GOLD, 0.45),
        ("Preferences", ["Travel style & pace", "Preferred currencies", "Climate / culture", "Must-haves & avoid"], TEAL, 3.5),
        ("Policy", ["Hard budget caps", "Date / duration windows", "Visa / compliance", "Corporate / family rules"], CORAL, 6.55),
    ]:
        add_card(slide, x, 1.3, 2.9, 2.55, CARD, col)
        strip = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x), Inches(1.3), Inches(2.9), Inches(0.1))
        strip.fill.solid()
        strip.fill.fore_color.rgb = col
        strip.line.fill.background()
        hd = slide.shapes.add_textbox(Inches(x + 0.15), Inches(1.5), Inches(2.6), Inches(0.35))
        _font(hd.text_frame.paragraphs[0], title, 14, bold=True, color=col, align=PP_ALIGN.CENTER)
        iy = 1.95
        for item in items:
            tb = slide.shapes.add_textbox(Inches(x + 0.2), Inches(iy), Inches(2.5), Inches(0.35))
            _font(tb.text_frame.paragraphs[0], f"•  {item}", 11, color=WHITE)
            iy += 0.4

    emotions = [
        ("Mood chips", "DEMO", GREEN),
        ("Text sentiment", "PARTIAL", GOLD),
        ("Selfie expression", "FUTURE", SKY),
        ("Voice emotion", "FUTURE", CORAL),
    ]
    ex = 0.45
    for label, scope, col in emotions:
        add_card(slide, ex, 4.1, 2.25, 0.95, CARD, col)
        tb = slide.shapes.add_textbox(Inches(ex + 0.1), Inches(4.2), Inches(2.05), Inches(0.75))
        _font(tb.text_frame.paragraphs[0], label, 12, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        p = tb.text_frame.add_paragraph()
        _font(p, scope, 10, bold=True, color=col, align=PP_ALIGN.CENTER)
        ex += 2.35

    add_callout_pro(
        slide,
        "Emotion is not only chips — roadmap includes CV selfie, voice ML, and sentiment over free-text intent.",
        y=5.3,
        accent=GOLD,
    )
    note = slide.shapes.add_textbox(Inches(0.55), Inches(6.2), Inches(8.9), Inches(0.5))
    _font(
        note.text_frame.paragraphs[0],
        "Future context: RAG over destination knowledge + social/friend-circle signals feeding the same 3 Ps layer.",
        11,
        color=SILVER,
    )
    add_footer(slide, num, "Architecture")


def slide_demo_vs_future(prs):
    slide = new_slide(prs)
    add_bg_decor(slide, GREEN)
    add_header(slide, "Demo In Scope vs Future", "Clear boundaries for judges — no ambiguity", GREEN)
    num = _n()

    if try_add_picture(slide, DEMO_PNG, 0.4, 1.25, 9.2, 4.5):
        add_footer(slide, num, "Architecture")
        return

    # Dense fallback mirroring Mermaid demo_vs_future.mmd
    add_card(slide, 0.4, 1.3, 4.5, 4.5, CARD, GREEN)
    hd = slide.shapes.add_textbox(Inches(0.55), Inches(1.4), Inches(4.2), Inches(0.4))
    _font(hd.text_frame.paragraphs[0], "DEMO MVP — IN SCOPE", 14, bold=True, color=GREEN)
    demo_items = [
        "Manual mood chips + free-text intent",
        "3 Ps lite from form (Profile · Prefs · Policy)",
        "TravelAgent: Emotional · Matcher · Cluster",
        "Gemini 2.5 + automatic mock fallback",
        "Concepts · inspiration boards · match scores",
        "Multi-currency cost guidance",
        "FastAPI + Web UI live end-to-end demo",
    ]
    y = 1.95
    for item in demo_items:
        tb = slide.shapes.add_textbox(Inches(0.65), Inches(y), Inches(4.0), Inches(0.4))
        _font(tb.text_frame.paragraphs[0], f"✓  {item}", 11, color=WHITE)
        y += 0.48

    add_card(slide, 5.1, 1.3, 4.5, 4.5, CARD, CORAL)
    hd2 = slide.shapes.add_textbox(Inches(5.25), Inches(1.4), Inches(4.2), Inches(0.4))
    _font(hd2.text_frame.paragraphs[0], "FUTURE / ROADMAP", 14, bold=True, color=CORAL)
    future_items = [
        "Selfie facial expression (real CV/ML)",
        "Voice emotion recognition",
        "RAG over destination knowledge base",
        "Social / friend-circle graph context",
        "Booking API orchestration",
        "Interactive maps & AI imagery",
        "Accounts, saved boards, mobile PWA",
    ]
    y = 1.95
    for item in future_items:
        tb = slide.shapes.add_textbox(Inches(5.35), Inches(y), Inches(4.0), Inches(0.4))
        _font(tb.text_frame.paragraphs[0], f"→  {item}", 11, color=SILVER)
        y += 0.48

    add_callout_pro(
        slide,
        "Source diagrams: docs/architecture/demo_vs_future.mmd · agent_architecture.mmd",
        y=6.0,
        accent=MUTED,
    )
    add_footer(slide, num, "Architecture")


def slide_capabilities(prs):
    slide = new_slide(prs)
    add_bg_decor(slide)
    add_header(slide, "Core Capabilities", "Four pillars of the AI Travel Inspirator")
    num = _n()

    pillars = [
        ("Mood & Multi-Modal\nEmotion", "Chips today; selfie/voice/sentiment roadmap", GOLD, 0.55),
        ("3 Ps Context\nEngine", "Profile · Preferences · Policy enrichment", TEAL, 2.75),
        ("Explainable\nInspiration", "Match scores + why-it-fits reasoning", CORAL, 4.95),
        ("Global Multi-\nCurrency", "USD · EUR · GBP · INR · AED · JPY · AUD", SKY, 7.15),
    ]
    for title, desc, col, x in pillars:
        add_card(slide, x, 1.4, 2.05, 2.15, CARD, col)
        tb = slide.shapes.add_textbox(Inches(x + 0.1), Inches(1.55), Inches(1.85), Inches(1.8))
        _font(tb.text_frame.paragraphs[0], title, 11, bold=True, color=col, align=PP_ALIGN.CENTER)
        p2 = tb.text_frame.add_paragraph()
        _font(p2, desc, 9, color=SILVER, align=PP_ALIGN.CENTER)

    add_bullets_pro(
        slide,
        [
            "Captures emotion and intent, then overlays 3 Ps so recommendations respect real constraints.",
            "Generates curated travel concepts and visual inspiration boards — not raw search listings.",
            "Honest trade-offs on cost (multi-currency), duration, and emotional fit for each destination.",
        ],
        top=3.85,
        size=13,
    )
    add_footer(slide, num, "Product")


def slide_competitive_landscape(prs):
    slide = new_slide(prs)
    add_bg_decor(slide, GOLD)
    add_header(slide, "Competitive Landscape", "Named players — research-backed positioning", GOLD)
    num = _n()

    players = [
        ("Airbnb", "Emotion / discovery-led shopping — Tiny Homes, Wish Lists, belonging narrative; AI for experience, not only price", GOLD),
        ("Booking.com", "Functional utility — price, availability, conversion; AI for search efficiency", TEAL),
        ("Expedia / Hotels.com", "AI agents & connected trip orchestration across inventory", CORAL),
        ("Google Travel / Flights", "Destination research + powerful price tools — limited emotion layer", SKY),
        ("MakeMyTrip Myra", "GenAI trip assistant (India example) — planning help, not emotion-first boards", GREEN),
        ("Instagram / TikTok", "Massive inspiration without personalization, cost, or policy context", MUTED),
    ]
    y = 1.3
    for name, desc, col in players:
        add_card(slide, 0.45, y, 9.1, 0.7, CARD, col)
        nb = slide.shapes.add_textbox(Inches(0.6), Inches(y + 0.15), Inches(2.3), Inches(0.4))
        _font(nb.text_frame.paragraphs[0], name, 12, bold=True, color=col)
        db = slide.shapes.add_textbox(Inches(3.0), Inches(y + 0.12), Inches(6.35), Inches(0.5))
        tf = db.text_frame
        tf.word_wrap = True
        _font(tf.paragraphs[0], desc, 10, color=WHITE)
        y += 0.78

    add_callout_pro(
        slide,
        "Wedge: emotion + 3 Ps + multi-modal signals + explainable boards — BETWEEN social inspiration and OTA booking.",
        y=6.1,
        accent=TEAL,
    )
    add_footer(slide, num, "Market")


def slide_competitive_differentiation(prs):
    slide = new_slide(prs)
    add_bg_decor(slide, TEAL)
    add_header(slide, "Competitive Differentiation", "How we occupy the unmet gap", TEAL)
    num = _n()

    headers = [("Criteria", 0.4, 1.9), ("Typical Players", 2.4, 3.3), ("AI Travel Inspirator", 5.8, 3.8)]
    for label, x, w in headers:
        add_box_text(slide, x, 1.3, w, 0.4, label, CARD_LIGHT, GOLD if x > 5 else SILVER, GOLD, 10, True)

    rows = [
        ("Primary lens", "Price / inventory / feeds", "Emotion + 3 Ps context"),
        ("Inspiration", "Social virality or categories", "Explainable mood boards"),
        ("Personalisation", "Generic profiles / filters", "Multi-modal + policy-aware"),
        ("Currency", "Local or single-market", "USD·EUR·GBP·INR·AED·JPY·AUD"),
        ("AI role", "Search, agents, conversion", "Emotional fit + reasoning"),
        ("Gap filled", "Inspire XOR book", "Bridge inspire → decide"),
    ]
    y = 1.8
    for i, (criteria, alt, ours) in enumerate(rows):
        fill = CARD if i % 2 == 0 else CARD_LIGHT
        c1 = add_card(slide, 0.4, y, 1.9, 0.55, fill, MUTED)
        _font(c1.text_frame.paragraphs[0], criteria, 9, bold=True, color=SILVER, align=PP_ALIGN.CENTER)
        c2 = add_card(slide, 2.4, y, 3.3, 0.55, fill, MUTED)
        _font(c2.text_frame.paragraphs[0], alt, 9, color=MUTED, align=PP_ALIGN.CENTER)
        c3 = add_card(slide, 5.8, y, 3.8, 0.55, fill, TEAL)
        _font(c3.text_frame.paragraphs[0], ours, 9, bold=True, color=TEAL, align=PP_ALIGN.CENTER)
        y += 0.6

    add_callout_pro(
        slide,
        "We do not compete on more destinations — we compete on better emotional + contextual fit.",
        y=5.7,
        accent=GOLD,
    )
    add_footer(slide, num, "Market")


def slide_mvp_scope(prs):
    slide = new_slide(prs)
    add_bg_decor(slide, GREEN)
    add_header(slide, "Hackathon MVP Scope", "Demo in scope vs deferred — aligned to agent roadmap", GREEN)
    num = _n()

    cols = [
        ("IN SCOPE (DEMO)", "Must deliver", GREEN, [
            "Mood chips + intent form", "3 Ps lite from inputs", "3-engine TravelAgent",
            "Gemini + mock fallback", "Concepts + boards", "Multi-currency costs",
            "FastAPI + Web UI", "Live E2E demo",
        ]),
        ("PARTIAL", "Demo-ready", GOLD, [
            "Text sentiment signals", "Match scoring 0–100%", "Sample itineraries",
            "Colour mood palettes", "Multi-provider LLM hooks",
        ]),
        ("FUTURE", "Post-hackathon", CORAL, [
            "Selfie CV / voice ML", "RAG destination KB", "Social graph context",
            "Booking integration", "Maps & AI imagery", "Accounts / mobile PWA",
        ]),
    ]
    cx = 0.45
    for title, sub, col, items in cols:
        add_card(slide, cx, 1.3, 2.95, 0.55, col, col)
        hd = slide.shapes.add_textbox(Inches(cx), Inches(1.37), Inches(2.95), Inches(0.4))
        _font(hd.text_frame.paragraphs[0], title, 11, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
        sb = slide.shapes.add_textbox(Inches(cx), Inches(1.9), Inches(2.95), Inches(0.25))
        _font(sb.text_frame.paragraphs[0], sub, 8, color=MUTED, align=PP_ALIGN.CENTER)
        add_card(slide, cx, 2.2, 2.95, 3.7, CARD, col)
        y = 2.35
        for item in items:
            tb = slide.shapes.add_textbox(Inches(cx + 0.15), Inches(y), Inches(2.65), Inches(0.35))
            _font(tb.text_frame.paragraphs[0], f"•  {item}", 10, color=WHITE)
            y += 0.4
        cx += 3.15

    add_callout_pro(
        slide,
        "MVP Focus: Prove emotion-first discovery with 3 Ps context — mood to actionable inspiration in < 30 seconds.",
        accent=GREEN,
    )
    add_footer(slide, num, "MVP Scope")


def slide_mvp_deliverables(prs):
    slide = new_slide(prs)
    add_bg_decor(slide)
    add_header(slide, "MVP Deliverables & Success Criteria", "Concrete outputs judges can evaluate")
    num = _n()

    deliverables = [
        ("Web Application", "Mood picker, multi-currency form, results renderer", GOLD, 0.55, 1.35),
        ("AI Agent Backend", "TravelAgent, 3 engines, Gemini + mock, 3 Ps context", TEAL, 5.1, 1.35),
        ("REST API", "POST /api/inspire — typed JSON contract", CORAL, 0.55, 3.55),
        ("Demo Flow", "Emotion → context → concepts → boards → costs", SKY, 5.1, 3.55),
    ]
    for title, desc, col, x, y in deliverables:
        add_card(slide, x, y, 4.25, 1.85, CARD, col)
        strip = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x), Inches(y), Inches(4.25), Inches(0.1))
        strip.fill.solid()
        strip.fill.fore_color.rgb = col
        strip.line.fill.background()
        tb = slide.shapes.add_textbox(Inches(x + 0.2), Inches(y + 0.25), Inches(3.85), Inches(1.4))
        _font(tb.text_frame.paragraphs[0], title, 14, bold=True, color=col)
        p2 = tb.text_frame.add_paragraph()
        _font(p2, desc, 11, color=SILVER)

    add_card(slide, 0.55, 5.55, 8.9, 0.95, CARD_LIGHT, TEAL)
    sc = slide.shapes.add_textbox(Inches(0.75), Inches(5.65), Inches(8.5), Inches(0.75))
    _font(sc.text_frame.paragraphs[0], "Success Criteria", 12, bold=True, color=TEAL)
    p2 = sc.text_frame.add_paragraph()
    _font(
        p2,
        "Live demo without errors  ·  Personalised to emotion/budget/currency  ·  Agentic architecture clear  ·  README + .env",
        11,
        color=WHITE,
    )
    add_footer(slide, num, "MVP Scope")


def slide_user_inputs(prs):
    slide = new_slide(prs)
    add_bg_decor(slide)
    add_header(slide, "User Inputs & Signals", "Emotion + 3 Ps intake (global traveller)")
    num = _n()

    inputs = [
        ("Emotion", "Mood chips (DEMO) · free-text intent · sentiment · selfie/voice (FUTURE)", GOLD),
        ("Profile", "Party type, travel personality signals, optional departure city", TEAL),
        ("Preferences", "Travel style · climate · pace · preferred currency display", CORAL),
        ("Policy", "Budget band · duration · hard constraints (visa, accessibility)", SKY),
        ("Currencies", "USD · EUR · GBP · INR · AED · JPY · AUD — emotions cross borders", GREEN),
    ]
    y = 1.35
    for i, (label, desc, col) in enumerate(inputs):
        add_card(slide, 0.55, y, 8.9, 0.7, CARD if i % 2 == 0 else CARD_LIGHT, col)
        lb = slide.shapes.add_textbox(Inches(0.75), Inches(y + 0.15), Inches(1.8), Inches(0.4))
        _font(lb.text_frame.paragraphs[0], label, 12, bold=True, color=col)
        db = slide.shapes.add_textbox(Inches(2.7), Inches(y + 0.15), Inches(6.5), Inches(0.45))
        tf = db.text_frame
        tf.word_wrap = True
        _font(tf.paragraphs[0], desc, 11, color=WHITE)
        y += 0.8

    add_footer(slide, num, "Product")


def slide_ai_features(prs):
    slide = new_slide(prs)
    add_bg_decor(slide, TEAL)
    add_header(slide, "TravelAgent Engines", "Three-module recommendation pipeline", TEAL)
    num = _n()

    engines = [
        ("01", "Emotional Engine", "Emotion + intent + 3 Ps → emotional profile, personality, priorities", GOLD),
        ("02", "Destination Matcher", "Scores 0–100 on personality, budget/policy, emotional fit", TEAL),
        ("03", "Inspiration Cluster", "Thematic mood-board clusters with palettes & activities", CORAL),
    ]
    y = 1.35
    for num_str, title, desc, col in engines:
        add_card(slide, 0.55, y, 8.9, 0.85, CARD, col)
        nb = slide.shapes.add_textbox(Inches(0.75), Inches(y + 0.2), Inches(0.5), Inches(0.45))
        _font(nb.text_frame.paragraphs[0], num_str, 16, bold=True, color=col)
        tb = slide.shapes.add_textbox(Inches(1.35), Inches(y + 0.15), Inches(2.6), Inches(0.55))
        _font(tb.text_frame.paragraphs[0], title, 13, bold=True, color=WHITE)
        db = slide.shapes.add_textbox(Inches(4.1), Inches(y + 0.2), Inches(5.1), Inches(0.5))
        _font(db.text_frame.paragraphs[0], desc, 11, color=SILVER)
        y += 1.0

    add_bullets_pro(
        slide,
        [
            "LLM reasoning with Pydantic-validated JSON — reliable, typed API responses.",
            "Gemini 2.5 primary with automatic mock fallback for demo resilience.",
            "RAG retrieval plugs in ahead of engines (FUTURE) without changing UI contracts.",
        ],
        top=4.5,
        size=12,
    )
    add_footer(slide, num, "Architecture")


def slide_outputs(prs):
    slide = new_slide(prs)
    add_bg_decor(slide)
    add_header(slide, "Generated Outputs", "What the traveller receives")
    num = _n()

    for x, title, items, col in [
        (0.55, "Curated Travel Concepts", [
            "2–3 themed trip concepts with taglines",
            "Destinations with 0–100% match scores",
            "Emotional fit rationale per destination",
            "Multi-currency cost estimates & duration",
            "Sample itineraries & budget breakdown",
        ], GOLD),
        (5.05, "Personalised Inspiration Board", [
            "Emotional summary & tailored quote",
            "2–3 thematic mood-board clusters",
            "Visual mood descriptions",
            "Colour palettes (hex codes)",
            "Grouped destinations & activities",
        ], TEAL),
    ]:
        add_card(slide, x, 1.35, 4.35, 4.85, CARD, col)
        strip = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x), Inches(1.35), Inches(4.35), Inches(0.12))
        strip.fill.solid()
        strip.fill.fore_color.rgb = col
        strip.line.fill.background()
        hd = slide.shapes.add_textbox(Inches(x + 0.2), Inches(1.55), Inches(3.95), Inches(0.4))
        _font(hd.text_frame.paragraphs[0], title, 14, bold=True, color=col)
        iy = 2.05
        for item in items:
            tb = slide.shapes.add_textbox(Inches(x + 0.25), Inches(iy), Inches(3.85), Inches(0.35))
            _font(tb.text_frame.paragraphs[0], f"•  {item}", 11, color=WHITE)
            iy += 0.42

    add_footer(slide, num, "Product")


def slide_system_architecture(prs):
    """Primary architecture slide — PNG if available, else dense boxes mirroring Mermaid."""
    slide = new_slide(prs)
    add_bg_decor(slide)
    add_header(slide, "System Architecture", "Agentic stack — mirrors docs/architecture/agent_architecture.mmd")
    num = _n()

    if try_add_picture(slide, AGENT_PNG, 0.35, 1.2, 9.3, 5.5):
        add_footer(slide, num, "Architecture")
        return

    # Dense technical boxes mirroring Mermaid structure
    rows = [
        (1.25, "UI", TEAL, ["Web UI / Mood Form", "Multi-modal Capture (selfie · voice · text)"]),
        (2.15, "API", GOLD, ["FastAPI  ·  POST /inspire  ·  /health · /diagnose"]),
        (3.0, "Emotion", CORAL, ["Mood chips DEMO", "Sentiment PARTIAL", "Selfie/Voice FUTURE"]),
        (3.85, "Context", TEAL, ["Profile", "Preferences", "Policy", "Social graph FUTURE"]),
        (4.7, "RAG", MUTED, ["Destination KB · Reviews · Culture — FUTURE"]),
        (5.45, "Agent→LLM→Out", SKY, ["3 Engines", "Gemini 2.5 + Mock", "Concepts · Boards · Multi-currency"]),
    ]
    for y, label, col, boxes in rows:
        add_card(slide, 0.4, y, 1.35, 0.65, CARD_LIGHT, col)
        lb = slide.shapes.add_textbox(Inches(0.45), Inches(y + 0.15), Inches(1.25), Inches(0.4))
        _font(lb.text_frame.paragraphs[0], label, 10, bold=True, color=col, align=PP_ALIGN.CENTER)
        bx = 1.9
        bw = (7.7 / len(boxes)) - 0.08
        for b in boxes:
            add_box_text(slide, bx, y, bw, 0.65, b, CARD, WHITE, col, 9)
            bx += bw + 0.08

    add_footer(slide, num, "Architecture")


def slide_architecture_mermaid_note(prs):
    """Secondary slide with Mermaid source / diagram reference."""
    slide = new_slide(prs)
    add_bg_decor(slide, SKY)
    add_header(slide, "Architecture Diagram Source", "Mermaid definitions live in docs/architecture/", SKY)
    num = _n()

    mmd_path = ARCH_DIR / "agent_architecture.mmd"
    snippet = ""
    if mmd_path.exists():
        lines = mmd_path.read_text(encoding="utf-8").splitlines()
        # Skip theme init; show flowchart body
        body = [ln for ln in lines if not ln.startswith("%%")]
        snippet = "\n".join(body[:22])
    else:
        snippet = "flowchart TB\n  UI --> API --> Emotion --> Context --> Agent --> LLM --> Outputs"

    add_card(slide, 0.4, 1.3, 9.2, 4.9, CARD, SKY)
    tb = slide.shapes.add_textbox(Inches(0.55), Inches(1.4), Inches(8.9), Inches(4.6))
    tf = tb.text_frame
    tf.word_wrap = True
    first = True
    for line in snippet.splitlines():
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        _font(p, line[:95], 9, color=SILVER, name=FONT_MONO)

    note = slide.shapes.add_textbox(Inches(0.5), Inches(6.35), Inches(9), Inches(0.4))
    _font(
        note.text_frame.paragraphs[0],
        "Files: agent_architecture.mmd · demo_vs_future.mmd  |  Render: npx @mermaid-js/mermaid-cli@11 …",
        10,
        color=MUTED,
    )
    add_footer(slide, num, "Architecture")


def slide_data_flow(prs):
    slide = new_slide(prs)
    add_bg_decor(slide)
    add_header(slide, "Data Flow", "End-to-end request lifecycle with 3 Ps + emotion")
    num = _n()

    steps = [
        ("User", "Emotion", GOLD),
        ("Context", "3 Ps", TEAL),
        ("API", "Validate", CORAL),
        ("Agent", "Engines", SKY),
        ("LLM", "Gemini", GOLD),
        ("UI", "Boards", GREEN),
    ]
    px = 0.35
    for title, sub, col in steps:
        add_box_text(slide, px, 1.45, 1.4, 0.95, f"{title}\n{sub}", CARD, col, col, 10, True)
        if px < 7.5:
            add_arrow(slide, px + 1.45, 1.9, px + 1.7, 1.9, TEAL)
        px += 1.55

    add_box_text(
        slide, 1.5, 2.7, 7.0, 0.5,
        "Optional FUTURE: RAG retrieve + social graph → inject into Context before Agent",
        CARD_LIGHT, MUTED, MUTED, 11, True,
    )

    ox = 0.7
    for title, col in [("Profile", GOLD), ("Concepts", TEAL), ("Board", CORAL), ("Multi-FX Costs", SKY)]:
        add_box_text(slide, ox, 3.5, 2.0, 0.75, title, CARD, col, col, 11, True)
        ox += 2.25

    add_bullets_pro(
        slide,
        [
            "Pydantic validation guarantees typed TravelResponse for the UI.",
            "Mock LLM path preserves demo if Gemini is unreachable — same JSON contract.",
        ],
        top=4.55,
        size=12,
    )
    add_footer(slide, num, "Architecture")


def slide_tech_stack(prs):
    slide = new_slide(prs)
    add_bg_decor(slide)
    add_header(slide, "Technology Stack", "MVP choices — Gemini, multi-currency, agentic ready")
    num = _n()

    for title, items, col, x in [
        ("Backend", ["Python 3", "FastAPI", "Uvicorn", "Pydantic v2", "httpx"], GOLD, 0.55),
        ("Frontend", ["HTML5 / CSS3", "JavaScript", "Multi-currency UI", "Responsive"], TEAL, 3.55),
        ("AI / Context", ["Gemini 2.5", "Mock fallback", "3 Ps context", "RAG hooks (future)"], CORAL, 6.55),
    ]:
        add_card(slide, x, 1.35, 2.75, 0.5, col, col)
        _font(slide.shapes[-1].text_frame.paragraphs[0], title, 13, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
        y = 2.0
        for item in items:
            add_box_text(slide, x, y, 2.75, 0.42, item, CARD, WHITE, col, 10)
            y += 0.5

    add_bullets_pro(
        slide,
        [
            "Environment-driven config via .env — provider, API keys, model, currency defaults.",
            "Modular agent design supports post-hackathon RAG, CV/voice emotion, maps, booking.",
        ],
        top=4.85,
        size=12,
    )
    add_footer(slide, num, "Architecture")


def slide_demo_flow(prs):
    slide = new_slide(prs)
    add_bg_decor(slide, GOLD)
    add_header(slide, "Live Demo Script", "Recommended flow for judges", GOLD)
    num = _n()

    steps = [
        ("1", "Launch App", "Open web UI — emotion-first interface"),
        ("2", "Enter Signals", "Peaceful · Solo · USD/EUR/INR budget band"),
        ("3", "Submit", "Emotion + 3 Ps → TravelAgent pipeline"),
        ("4", "Review Profile", "Travel personality & priorities"),
        ("5", "Explore Concepts", "Match scores & multi-currency costs"),
        ("6", "Browse Board", "Thematic clusters & colour palettes"),
    ]
    for i, (num_s, title, desc) in enumerate(steps):
        row, col = i // 2, i % 2
        x, y = 0.55 + col * 4.75, 1.4 + row * 1.75
        badge = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(0.5), Inches(0.5))
        badge.fill.solid()
        badge.fill.fore_color.rgb = GOLD
        badge.line.fill.background()
        cn = slide.shapes.add_textbox(Inches(x), Inches(y + 0.08), Inches(0.5), Inches(0.35))
        _font(cn.text_frame.paragraphs[0], num_s, 14, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
        add_card(slide, x + 0.6, y, 3.95, 0.95, CARD, GOLD)
        tb = slide.shapes.add_textbox(Inches(x + 0.75), Inches(y + 0.1), Inches(3.65), Inches(0.75))
        _font(tb.text_frame.paragraphs[0], title, 12, bold=True, color=GOLD)
        p2 = tb.text_frame.add_paragraph()
        _font(p2, desc, 10, color=SILVER)

    add_footer(slide, num, "Demo")


def slide_roadmap(prs):
    slide = new_slide(prs)
    add_bg_decor(slide)
    add_header(slide, "Post-Hackathon Roadmap", "Scaling the agentic stack beyond MVP")
    num = _n()

    phases = [
        ("Phase 2", "Multi-modal emotion — selfie CV expression + voice emotion ML"),
        ("Phase 3", "RAG over destination knowledge, reviews, culture & seasonality"),
        ("Phase 4", "Social / friend-circle context feeding the 3 Ps layer"),
        ("Phase 5", "Real-time pricing & booking API orchestration"),
        ("Phase 6", "Maps, AI imagery, accounts, saved boards, mobile PWA"),
    ]
    y = 1.4
    for i, (phase, desc) in enumerate(phases):
        col = [GOLD, TEAL, CORAL, SKY, GOLD][i]
        add_card(slide, 0.55, y, 8.9, 0.7, CARD if i % 2 == 0 else CARD_LIGHT, col)
        tb = slide.shapes.add_textbox(Inches(0.75), Inches(y + 0.1), Inches(8.5), Inches(0.5))
        _font(tb.text_frame.paragraphs[0], phase, 12, bold=True, color=col)
        p2 = tb.text_frame.add_paragraph()
        _font(p2, desc, 11, color=WHITE)
        y += 0.8

    add_callout_pro(
        slide,
        "MVP validates the emotion→inspire hypothesis; roadmap scales context, RAG, and booking.",
        accent=TEAL,
        y=5.7,
    )
    add_footer(slide, num, "Roadmap")


def slide_thank_you(prs):
    slide = new_slide(prs)
    add_bg_decor(slide, GOLD)

    banner = add_card(slide, 1.25, 1.6, 7.5, 2.2, CARD_LIGHT, GOLD)
    tb = slide.shapes.add_textbox(Inches(1.5), Inches(1.85), Inches(7.0), Inches(0.9))
    _font(tb.text_frame.paragraphs[0], "Thank You", 44, bold=True, color=GOLD, align=PP_ALIGN.CENTER)

    sub = slide.shapes.add_textbox(Inches(1.5), Inches(2.75), Inches(7.0), Inches(0.9))
    _font(sub.text_frame.paragraphs[0], "AI Travel Inspirator", 20, color=WHITE, align=PP_ALIGN.CENTER)
    p2 = sub.text_frame.add_paragraph()
    _font(p2, "Discover where your heart wants to go — anywhere in the world", 13, color=SILVER, align=PP_ALIGN.CENTER)

    gold_line = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(3.5), Inches(4.15), Inches(3.0), Inches(0.05))
    gold_line.fill.solid()
    gold_line.fill.fore_color.rgb = GOLD
    gold_line.line.fill.background()

    info = slide.shapes.add_textbox(Inches(1), Inches(4.45), Inches(8), Inches(1.5))
    _font(info.text_frame.paragraphs[0], f"Team {TEAM}  ·  {ORG}", 14, bold=True, color=TEAL, align=PP_ALIGN.CENTER)
    for line in ["GitHub: AviCoooool/AITravel", "Demo: http://localhost:8002", "We welcome your questions"]:
        p = info.text_frame.add_paragraph()
        _font(p, line, 12, color=MUTED, align=PP_ALIGN.CENTER)

    add_footer(slide, _n())


def build_presentation() -> Path:
    SLIDE_NUM[0] = 0
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    ARCH_DIR.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    # Opening
    slide_title(prs)
    slide_team(prs)

    # 01 Opportunity
    slide_section(prs, "01", "The Opportunity", "Market scale, problem, and underserved emotion layer", CORAL)
    slide_executive_summary(prs)
    slide_market_opportunity(prs)
    slide_problem(prs)

    # 02 Solution
    slide_section(prs, "02", "Our Solution", "Emotion-first agent with 3 Ps context", TEAL)
    slide_solution_proposal(prs)
    slide_agentic_architecture_overview(prs)
    slide_three_ps_and_emotion(prs)
    slide_demo_vs_future(prs)
    slide_capabilities(prs)

    # 02b Competition
    slide_section(prs, "02b", "Market Position", "Named competitors and our wedge", GOLD)
    slide_competitive_landscape(prs)
    slide_competitive_differentiation(prs)

    # 03 MVP
    slide_section(prs, "03", "Hackathon MVP", "Scope and deliverables", GREEN)
    slide_mvp_scope(prs)
    slide_mvp_deliverables(prs)

    # 04 Architecture & Demo
    slide_section(prs, "04", "Architecture & Demo", "Technical deep dive and live demonstration", SKY)
    slide_user_inputs(prs)
    slide_ai_features(prs)
    slide_outputs(prs)
    slide_system_architecture(prs)
    slide_architecture_mermaid_note(prs)
    slide_data_flow(prs)
    slide_tech_stack(prs)
    slide_demo_flow(prs)

    # 05 Close
    slide_section(prs, "05", "What's Next", "Roadmap and closing", GOLD)
    slide_roadmap(prs)
    slide_thank_you(prs)

    prs.save(str(OUTPUT))
    return OUTPUT


if __name__ == "__main__":
    path = build_presentation()
    print(f"Presentation saved: {path} ({len(Presentation(path).slides)} slides)")
    print(f"Mermaid PNG agent_architecture: {'YES' if AGENT_PNG.exists() else 'NO'}")
    print(f"Mermaid PNG demo_vs_future: {'YES' if DEMO_PNG.exists() else 'NO'}")
